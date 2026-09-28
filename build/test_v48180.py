from pathlib import Path
import base64, os, shutil, sqlite3, tempfile, threading
from concurrent.futures import ThreadPoolExecutor
import backend

ROOT=Path(__file__).resolve().parent
assert backend.APP_VERSION=='4.8.18.0-DESKTOP', backend.APP_VERSION
base=Path(tempfile.mkdtemp(prefix='sokhna_v48180_'))
root=base/'Sokhna Port Data'
s=backend.DesktopStore(root)
try:
    top={p.name for p in root.iterdir() if p.is_dir()}
    assert top=={'Database','Attachments','Backup Recovery','Excel','Mobile SQL Exchange','Logs'},top
    assert s.storage_audit()['quickCheck']=='ok'
    now=backend.now_iso()
    s.put('companies',{'id':'c1','name':'شركة اختبار 4.8.18','companySerial':1,'createdAt':now,'modifiedAt':now})
    s.put('items',{'id':'i1','companyId':'c1','name':'رخصة التشغيل','applicable':True,'position':1,'expiryDate':'2027-01-01','createdAt':now,'modifiedAt':now})
    s.put('inspections',{'id':'in1','companyId':'c1','date':'2026-09-29','inspectionCode':'INS-1-2026-1','inspectionNo':'INS-1-2026-1','createdAt':now,'modifiedAt':now})
    s.put('findings',{'id':'f1','inspectionId':'in1','companyId':'c1','no':'1','description':'ملاحظة اختبار','requirement':'متطلب','status':'Open','targetDate':'2026-10-01','createdAt':now,'modifiedAt':now})
    payload=base64.b64encode(b'attachment-v48180').decode('ascii')
    s.put('attachments',{'id':'a1','companyId':'c1','entityType':'item','entityId':'i1','filename':'proof.txt','mime':'text/plain','dataUrl':'data:text/plain;base64,'+payload,'createdAt':now,'modifiedAt':now})
    def put_one(i):
        t=backend.now_iso();s.put('items',{'id':f'ci{i}','companyId':'c1','name':f'بند ضغط {i}','applicable':True,'position':i+10,'expiryDate':'2027-12-31','createdAt':t,'modifiedAt':t})
    with ThreadPoolExecutor(max_workers=8) as ex:list(ex.map(put_one,range(40)))
    assert len([x for x in s.get_all('items') if str(x.get('id','')).startswith('ci')])==40
    assert s.storage_audit()['quickCheck']=='ok'
    live=Path(s.refresh_live_excel_now()); audit=s.validate_live_excel(live)
    assert live.is_dir() and audit['verifiedAgainstSQLite'] and audit['excelFiles']==1
    xlsx=Path(audit['files'][0]['path']); assert xlsx.is_file() and xlsx.stat().st_size>1024
    from openpyxl import load_workbook
    wb=load_workbook(xlsx,read_only=True,data_only=False)
    assert wb.sheetnames==['البنود','التفتيشات']
    assert wb['البنود']['B3'].value=='رخصة التشغيل'
    assert wb['التفتيشات']['G3'].value=='ملاحظة اختبار'
    wb.close()
    backup=Path(s.create_backup('MANUAL')); assert backup.is_file()
    ver=s.verify_backup(str(backup),True); assert ver['valid']
    c=s.get_one('companies','c1');s.put('companies',{**c,'name':'اسم مؤقت','modifiedAt':backend.now_iso()})
    assert s.get_one('companies','c1')['name']=='اسم مؤقت'
    restored=s.restore_backup(str(backup)); assert Path(restored['beforeRestore']).is_file()
    assert s.get_one('companies','c1')['name']=='شركة اختبار 4.8.18'
    assert s.storage_audit()['quickCheck']=='ok'
    for include in (False,True):
        p=Path(s.create_mobile_sqlite('',include)); assert p.is_file()
        con=sqlite3.connect(str(p))
        try:
            assert con.execute('select count(*) from records').fetchone()[0]>=4
            if include: assert con.execute('select count(*) from attachment_blobs').fetchone()[0]>=1
        finally: con.close()
    assert s.set_retention_days(45)==45 and s.get_retention_days()==45
    idx=s.refresh_attachment_index(); assert idx['indexed']>=1 and idx['missing']==0 and Path(idx['path']).is_file()
finally:
    s.close()

s2=backend.DesktopStore(root)
try:
    assert s2.storage_audit()['quickCheck']=='ok'
    assert s2.get_one('companies','c1')['name']=='شركة اختبار 4.8.18'
finally:s2.close()

js=(ROOT/'app'/'v4179_patch.js').read_text(encoding='utf-8')
assert "new MutationObserver(()=>{if($('view-sync')?.classList.contains('active'))sync79()}" not in js
assert 'if(next!==current)el.textContent=next' in js
assert 'renderSyncBefore79' in js and 'runSync79' in js
assert '#openInspectionAnalyticsBtn,#inspectionAnalyticsBtn' in js
assert ".observe(inspectionDetail79,{subtree:true,childList:true})" in js
html=(ROOT/'app'/'index.html').read_text(encoding='utf-8')
assert "const APP_VERSION='4.8.18.0-DESKTOP'" in html

class FakeServer:
    def __init__(self): self.stopped=threading.Event()
    def shutdown(self): self.stopped.set()
os.environ['SOKHNA_IDLE_SHUTDOWN_SECONDS']='1'
backend.note_heartbeat();fake=FakeServer();backend.start_idle_shutdown(fake)
assert fake.stopped.wait(2.5),'idle shutdown did not recover stale instance'
os.environ.pop('SOKHNA_IDLE_SHUTDOWN_SECONDS',None)

shutil.rmtree(base,ignore_errors=True)
print('v4.8.18.0 CORE/STORAGE/EXCEL/BACKUP/MOBILE/CONCURRENCY/RECOVERY QA PASS')
