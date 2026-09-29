from pathlib import Path
import hashlib, tempfile, time
import backend

ROOT=Path(__file__).resolve().parent
assert backend.APP_VERSION=='4.8.18.2-DESKTOP', backend.APP_VERSION

template=ROOT/'app'/'sugar_template.xlsx'
assert template.is_file()
assert hashlib.sha256(template.read_bytes()).hexdigest()=='55e2257b0362d58b62403816ea44bf02e7cfe0e643a189ff74b7b0ad0b11874d'

base=Path(tempfile.mkdtemp(prefix='sokhna_live_48182_'))
s=backend.DesktopStore(base/'Sokhna Port Data')
try:
    now=backend.now_iso()
    s.put('companies',{'id':'c1','name':'شركة السكر','companySerial':1,'createdAt':now,'modifiedAt':now})
    s.put('items',{'id':'i1','companyId':'c1','name':'رخصة التشغيل','issueDate':'2026-01-01','expiryDate':'2027-01-01','currentStatus':'سليم','applicable':True,'position':1,'createdAt':now,'modifiedAt':now})
    s.put('inspections',{'id':'in1','companyId':'c1','date':'2026-09-29','inspectionCode':'INS-1-2026-1','createdAt':now,'modifiedAt':now})
    s.put('findings',{'id':'f1','inspectionId':'in1','companyId':'c1','no':'1','requirement':'متطلب اختبار','referenceNo':'R1','description':'ملاحظة اختبار','status':'Open','targetDate':'2026-10-15','createdAt':now,'modifiedAt':now})
    live=Path(s.refresh_live_excel_now())
    assert live==Path(s.paths['excel_default'])
    files=sorted(live.rglob('*.xlsx'))
    assert len(files)==1,files
    assert files[0].relative_to(live).parts==('شركة السكر','شركة السكر.xlsx'),files[0]
    audit=s.validate_live_excel(live)
    assert audit['verifiedAgainstSQLite'] and audit['companyCount']==1 and audit['excelFiles']==1,audit
    from openpyxl import load_workbook
    wb=load_workbook(files[0],read_only=True,data_only=False)
    assert wb.sheetnames==['البنود','التفتيشات']
    assert wb['البنود']['B3'].value=='رخصة التشغيل'
    assert wb['التفتيشات']['G3'].value=='ملاحظة اختبار'
    wb.close()
    old_mtime=files[0].stat().st_mtime
    row=s.get_one('items','i1');row['currentStatus']='تم التحديث تلقائيًا';row['modifiedAt']=backend.now_iso();s.put('items',row)
    deadline=time.time()+6
    while time.time()<deadline and files[0].stat().st_mtime<=old_mtime:time.sleep(.15)
    assert files[0].stat().st_mtime>old_mtime,'Live workbook did not auto-update after store write'
finally:
    s.close()

be=(ROOT/'v4179_backend_ext.py').read_text(encoding='utf-8')
assert "/api/open-live-excel-folder" in be
assert "os.startfile(str(root.resolve()))" in be and "explorer.exe" in be
js=(ROOT/'app'/'v48182_live_excel.js').read_text(encoding='utf-8')
for token in ('فتح فولدر Live Excel','تحديث Live Excel الآن','ورقتين فقط','البنود','التفتيشات','/api/open-live-excel-folder'):
    assert token in js,token
assert "obsolete.remove()" in js
print('v4.8.18.2 LIVE EXCEL STRUCTURE/AUTO-UPDATE/OPEN-FOLDER QA PASS')
